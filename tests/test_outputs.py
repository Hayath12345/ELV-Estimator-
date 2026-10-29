"""Output generation: every file is produced, PDFs open, and the Excel formulas — when
recalculated by LibreOffice (if installed) — have no errors and reproduce the engine
totals to QAR 0.00."""
import shutil
import subprocess
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest

from elv_engine import pipeline
from elv_engine.outputs import package

COMPANY = {"name": "Test ELV Contracting W.L.L.", "address": "P.O. Box 1, Doha, Qatar", "phone": "+974 0000 0000"}


@pytest.fixture()
def generated(tmp_path, sample_project, priced_library, sample_dxfs):
    folder = tmp_path / "proj"
    (folder / "cache").mkdir(parents=True)
    (folder / "drawings").mkdir()
    # background linework for the layout drawings
    from elv_engine.extract import extract_drawing
    for d, f in zip(sample_project["drawings"], sample_dxfs):
        ex = extract_drawing(f, str(folder / "cache"), d["id"])
        d["extraction"] = ex
    sample_project["revision"]["current"] = 1
    sample_project["signoff"]["prepared"] = {"name": "A. Estimator", "role": "ELV Estimator", "at": "2026-10-28T10:00:00"}
    res = pipeline.compute(sample_project, priced_library.as_map(), pipeline.effective_fx(priced_library.fx(), {}))
    sel = {"full", "per", "summary", "layout", "racks", "cost", "zip"}
    out = package.generate(sample_project, res, folder, sel, COMPANY, {"name": "Tester"})
    return out, res


def test_all_outputs_generated(generated):
    import pymupdf
    out, res = generated
    folder = Path(out["folder"])
    names = out["files"]
    assert any(n.endswith("_Quotation.pdf") and "_CCTV_" not in n for n in names)
    assert any(n.endswith("_Quotation.xlsx") for n in names)
    assert sum(1 for n in names if n.endswith("_Quotation.pdf")) == 1 + len(res.selected)
    for n in ("_Combined_Summary.pdf", "_Layouts.pdf", "_Rack_Elevations.pdf", "_Internal_Costing.xlsx", "_Package.zip"):
        assert any(x.endswith(n) for x in names), n
    for n in names:
        if n.endswith(".pdf"):
            with pymupdf.open(folder / n) as doc:
                assert doc.page_count >= 1
    with pymupdf.open(next(folder / n for n in names if n.endswith("_Layouts.pdf"))) as doc:
        assert doc.page_count == len(res.floors) * (1 + len(res.selected))
        text = doc[0].get_text()
        assert "1:" in text and "LEGEND" in text
    with pymupdf.open(next(folder / n for n in names if n.endswith("_Quotation.pdf") and n.count("_") == 2)) as doc:
        text = "".join(p.get_text() for p in doc)
        from elv_engine.money import fmt
        assert fmt(res.boq["totals"]["base"]) in text
        assert "Qatari Riyals" in text and "Bill of quantities" in text and "Exclusions" in text
    z = zipfile.ZipFile(out["zip"])
    assert not any("Internal_Costing" in n for n in z.namelist()), "internal costing must not go to the client"
    assert (folder / "issue.json").exists()


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_excel_formulas_recalculate_without_errors(generated, tmp_path):
    import openpyxl
    out, res = generated
    folder = Path(out["folder"])
    conv = tmp_path / "lo"
    conv.mkdir()
    files = [str(folder / n) for n in out["files"] if n.endswith(".xlsx")]
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    subprocess.run([exe, "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(conv), *files],
                   check=True, capture_output=True, timeout=180)
    for f in conv.iterdir():
        wb = openpyxl.load_workbook(f, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    assert not (isinstance(c.value, str) and c.value.startswith("#") and c.value.endswith(("!", "?", "A"))), \
                        f"{f.name} {ws.title}!{c.coordinate} = {c.value}"
    t = res.boq["totals"]
    q = openpyxl.load_workbook(next(conv.glob("*_Quotation.xlsx")), data_only=True)["Summary"]
    found = [r for r in q.iter_rows(values_only=True) if r[0] == "TOTAL BASE PRICE (QAR)"]
    assert found and Decimal(str(found[0][2])).quantize(Decimal("0.01")) == t["base"]
    c = openpyxl.load_workbook(next(conv.glob("*_Internal_Costing.xlsx")), data_only=True)["Summary"]
    for row in c.iter_rows(min_row=4, values_only=True):
        if row[0] and row[6] is not None:
            assert abs(float(row[6])) < 0.005, row
