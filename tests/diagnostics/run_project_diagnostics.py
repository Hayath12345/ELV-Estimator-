"""Test-only helper: run one project end to end through the real backend (app.backend.Backend via
app.api.dispatch — the same method calls the UI makes over the js_api bridge) and write the
evidence for test-report sections 2-7, 9 and 11.

    python tests/diagnostics/run_project_diagnostics.py OUT_DIR --label S
        (no --drawings: generates the built-in two-floor sample set into OUT_DIR/inputs/S/)
    python tests/diagnostics/run_project_diagnostics.py OUT_DIR --label B --drawings "C:/path/Basement.dwg"

Everything runs in OUT_DIR/work (own settings, price library and projects folder), so the
user's settings, price library and projects are not touched. Prices: the run first evaluates the
BOQ as shipped (empty library), then applies the deterministic test prices from tests/conftest.py
(TEST_PRICES) to the work library only and generates every output.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import re
import shutil
import statistics
import subprocess
import sys
import time
from collections import Counter, defaultdict
from contextlib import contextmanager
from decimal import Decimal
from pathlib import Path

os.environ["ELV_NO_OPEN"] = "1"
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import numpy as np  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402

from app.api import dispatch  # noqa: E402
from app.backend import Backend  # noqa: E402
from elv_engine import catalog  # noqa: E402
from elv_engine.boq import all_lines  # noqa: E402
from elv_engine.extract.classify import looks_like_room_name, looks_like_room_number  # noqa: E402
from elv_engine.layout import type_label  # noqa: E402
from elv_engine.layout.place import camera_range  # noqa: E402
from elv_engine.layout.routing import LIMITED  # noqa: E402
from elv_engine.money import fmt  # noqa: E402
from elv_engine.outputs.common import system_name  # noqa: E402
from elv_engine.params import merged  # noqa: E402
from elv_engine.systems import BY_KEY, SYSTEM_KEYS  # noqa: E402

TIMES: dict[str, float] = {}
WEDGE_HALF_DEG = math.degrees(0.62)
ACADVER = {"AC1009": "R12", "AC1012": "R13", "AC1014": "R14", "AC1015": "R2000", "AC1018": "R2004", "AC1021": "R2007",
           "AC1024": "R2010", "AC1027": "R2013", "AC1032": "R2018"}
BAD_WORDS = [("TBC", re.compile(r"\bTBC\b")), ("not found", re.compile(r"not found", re.I)),
             ("placeholder", re.compile(r"placeholder", re.I)), ("nan", re.compile(r"\bnan\b", re.I)),
             ("None", re.compile(r"\bNone\b"))]
PDF_TERMS = [("TBC", re.compile(r"\bTBC\b")), ("[QUOTING COMPANY", re.compile(re.escape("[QUOTING COMPANY"))),
             ("nan", re.compile(r"\bnan\b", re.I)), ("None", re.compile(r"\bNone\b")), ("#REF", re.compile("#REF")),
             ("lorem", re.compile("lorem", re.I))]
XL_ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "Err:")


@contextmanager
def timed(name):
    t0 = time.perf_counter()
    yield
    TIMES[name] = round(time.perf_counter() - t0, 3)


def jdump(path: Path, obj) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=1, ensure_ascii=False,
                               default=lambda o: str(o) if isinstance(o, Decimal) else repr(o)), encoding="utf-8")


class Client:
    def __init__(self, b: Backend, log):
        self.b, self.log = b, log

    def __call__(self, m, a=None):
        t0 = time.perf_counter()
        r = dispatch(self.b, m, a or {})
        self.log.write(f"{time.strftime('%H:%M:%S')} {m} {json.dumps(a or {}, default=str)[:160]} -> "
                       f"{'ok' if r['ok'] else 'ERROR ' + str(r['error'])} ({time.perf_counter() - t0:.3f}s)\n")
        if not r["ok"]:
            raise RuntimeError(f"{m}: {r['error']}")
        return r["data"]

    def wait(self, job, timeout=900):
        t0 = time.time()
        while time.time() - t0 < timeout:
            x = dispatch(self.b, "get_job", {"id": job["id"]})["data"]  # polled; not logged per call
            if x["done"]:
                if x["error"]:
                    raise RuntimeError(f"job {job['kind']}: {x['error']}")
                return x
            time.sleep(0.05)
        raise TimeoutError(job)


# ======================================================================= section 3 — drawings
def drawing_diag(src: Path, dxf_path: Path, ex: dict, unit_m: float, drawing: dict) -> tuple[dict, list[dict]]:
    import ezdxf
    from ezdxf import recover

    from elv_engine.extract.dxf import RE_DOOR_TAG, RE_LABEL_LAYER
    try:
        doc = ezdxf.readfile(str(dxf_path))
        read_mode = "ezdxf.readfile"
    except Exception:  # noqa: BLE001
        doc, _ = recover.readfile(str(dxf_path))
        read_mode = "ezdxf.recover"
    hdr = doc.header
    dwg_ver = None
    if src.suffix.lower() == ".dwg":
        dwg_ver = src.read_bytes()[:6].decode("ascii", "replace")
    msp = doc.modelspace()
    direct = defaultdict(Counter)
    for e in msp:
        direct[e.dxf.get("layer", "0")][e.dxftype()] += 1

    memo: dict[str, tuple[Counter, int]] = {}

    def block_counts(name: str, stack=()) -> tuple[Counter, int]:
        if name in memo:
            return memo[name]
        if name in stack or len(stack) > 16:
            return Counter(), 0
        blk = doc.blocks.get(name)
        c, depth = Counter(), 1
        if blk is not None:
            for e in blk:
                lay = e.dxf.get("layer", "0")
                c[(lay, e.dxftype())] += 1
                if e.dxftype() == "INSERT":
                    sub, d = block_counts(e.dxf.name, stack + (name,))
                    for (l2, t2), n in sub.items():
                        c[(lay if l2 == "0" else l2, t2)] += n
                    depth = max(depth, d + 1)
        memo[name] = (c, depth)
        return memo[name]

    expanded = defaultdict(Counter)
    max_depth = 0
    n_inserts = 0
    for e in msp:
        if e.dxftype() == "INSERT":
            n_inserts += 1
            sub, d = block_counts(e.dxf.name)
            max_depth = max(max_depth, d)
            ins_layer = e.dxf.get("layer", "0")
            for (lay, t), n in sub.items():
                expanded[ins_layer if lay == "0" else lay][t] += n
    xrefs = []
    for b in doc.blocks:
        blk = b.block
        if blk is not None and getattr(blk, "is_xref", False):
            xrefs.append({"name": b.name, "path": blk.dxf.get("xref_path", "")})

    # texts (model space, attributes, and texts inside blocks up to depth 4 in world coordinates)
    texts = []

    def add_text(e, where):
        t = e.dxftype()
        try:
            if t == "MTEXT":
                s = e.plain_text()
                pt = e.dxf.insert
            else:
                s = e.dxf.text
                use_align = t in ("TEXT", "ATTRIB") and (e.dxf.get("halign", 0) or e.dxf.get("valign", 0))
                pt = e.dxf.align_point if use_align and e.dxf.hasattr("align_point") else e.dxf.insert
        except Exception:  # noqa: BLE001
            return
        s = " ".join(str(s).split())
        if not s:
            return
        texts.append({"text": s, "type": t, "layer": e.dxf.get("layer", "0"), "where": where,
                      "x": float(pt[0]), "y": float(pt[1])})

    def from_insert(ins, depth):
        for a in getattr(ins, "attribs", []):
            add_text(a, f"ATTRIB of {ins.dxf.name}")
        if depth > 4:
            return
        try:
            ves = list(ins.virtual_entities())
        except Exception:  # noqa: BLE001
            return
        for v in ves:
            if v.dxftype() in ("TEXT", "MTEXT"):
                add_text(v, f"block {ins.dxf.name}")
            elif v.dxftype() == "INSERT":
                from_insert(v, depth + 1)
    for e in msp:
        if e.dxftype() in ("TEXT", "MTEXT"):
            add_text(e, "modelspace")
        elif e.dxftype() == "INSERT":
            from_insert(e, 1)
    for t in texts:
        t["x_m"], t["y_m"] = round(t["x"] * unit_m, 3), round(t["y"] * unit_m, 3)
        t["room_name_like"] = looks_like_room_name(t["text"])
        t["room_number_like"] = looks_like_room_number(t["text"])
        t["door_tag"] = bool(RE_DOOR_TAG.match(t["text"]))
        t["label_layer"] = bool(RE_LABEL_LAYER.search(t["layer"]))
    labels = [t for t in texts if not t["door_tag"] and (t["room_name_like"] or t["room_number_like"] or t["label_layer"])]
    dups = []
    by_text = defaultdict(list)
    for t in labels:
        by_text[t["text"].upper()].append(t)
    for k, lst in by_text.items():
        for i in range(len(lst)):
            for j in range(i + 1, len(lst)):
                dm = math.hypot(lst[i]["x_m"] - lst[j]["x_m"], lst[i]["y_m"] - lst[j]["y_m"])
                if dm <= 3.0:
                    dups.append({"text": k, "a": [lst[i]["x_m"], lst[i]["y_m"]], "b": [lst[j]["x_m"], lst[j]["y_m"]],
                                 "dist_m": round(dm, 2)})

    # boundaries: closed polylines for room source lookup
    closed = []
    for e in msp:
        if e.dxftype() == "LWPOLYLINE" and e.closed:
            pts = [(p[0], p[1]) for p in e.get_points("xy")]
        elif e.dxftype() == "POLYLINE" and e.is_closed:
            pts = [(v.dxf.location[0], v.dxf.location[1]) for v in e.vertices]
        else:
            continue
        if len(pts) >= 3:
            xs, ys = [p[0] for p in pts], [p[1] for p in pts]
            closed.append({"layer": e.dxf.get("layer", "0"), "type": e.dxftype(), "bbox": (min(xs), min(ys), max(xs), max(ys))})

    rooms_json = []
    for r in ex["rooms"]:
        xs, ys = [p[0] for p in r["poly"]], [p[1] for p in r["poly"]]
        bb = (min(xs), min(ys), max(xs), max(ys))
        lx, ly = r["label_pt"]
        lab = min(texts, key=lambda t: math.hypot(t["x"] - lx, t["y"] - ly), default=None)
        lab_ok = lab is not None and math.hypot(lab["x"] - lx, lab["y"] - ly) * unit_m < 0.5
        src_b = None
        if r["boundary"] == "drawn":
            tol = 0.002 / unit_m
            src_b = next((c for c in closed if all(abs(a - b) <= tol for a, b in zip(c["bbox"], bb))), None)
        from elv_engine import geometry as G
        c = G.centroid([tuple(p) for p in r["poly"]])
        rooms_json.append({
            "id": f"{drawing['id']}:{r['id']}", "rid": r["id"], "drawing": drawing["name"], "floor": drawing["floor"],
            "name": r["name"], "number": r.get("number", ""), "category": r["category"], "comms": r.get("comms"),
            "centroid_m": [round(c[0] * unit_m, 3), round(c[1] * unit_m, 3)], "area_m2": r["area_m2"],
            "has_boundary": r["boundary"] == "drawn", "boundary": r["boundary"], "conf": r["conf"],
            "label_source_entity": lab["type"] if lab_ok else None, "label_layer": lab["layer"] if lab_ok else None,
            "label_where": lab["where"] if lab_ok else None,
            "boundary_source_entity": src_b["type"] if src_b else None, "boundary_layer": src_b["layer"] if src_b else None,
            "duplicate_of": None,
        })
    for i, a in enumerate(rooms_json):
        for b in rooms_json[:i]:
            if b["name"].upper() == a["name"].upper() and math.hypot(a["centroid_m"][0] - b["centroid_m"][0],
                                                                  a["centroid_m"][1] - b["centroid_m"][1]) <= 3.0:
                a["duplicate_of"] = b["id"]
                break
    ext = ex["extents"]
    info = {
        "file": src.name, "size_bytes": src.stat().st_size, "kind": ex["kind"], "read_with": read_mode,
        "dwg_version_magic": dwg_ver, "dwg_version": ACADVER.get(dwg_ver or "", None),
        "dxf_ACADVER": hdr.get("$ACADVER"), "dxf_release": ACADVER.get(hdr.get("$ACADVER", ""), None),
        "INSUNITS": hdr.get("$INSUNITS"), "MEASUREMENT": hdr.get("$MEASUREMENT"),
        "EXTMIN_header": [round(v, 3) for v in hdr.get("$EXTMIN", (0, 0, 0))],
        "EXTMAX_header": [round(v, 3) for v in hdr.get("$EXTMAX", (0, 0, 0))],
        "extents_used_units": ext, "extents_used_m": [round(v * unit_m, 3) for v in ext],
        "size_m": [round((ext[2] - ext[0]) * unit_m, 2), round((ext[3] - ext[1]) * unit_m, 2)],
        "unit_m_used": unit_m, "scale_conf": ex.get("scale_conf"), "scale_note": ex.get("scale_note"),
        "plot_scale": ex.get("plot_scale"), "calibration": drawing.get("calibration"),
        "scale_derivation": ("calibration" if (drawing.get("calibration") or {}).get("unit_m") else
                             "infer_units($INSUNITS, robust extents, closed-polyline areas)"),
        "origin_used": "drawing WCS as-is (no translation/rotation; UCS and base point ignored)",
        "app_counts": ex.get("counts"), "app_warnings": ex.get("warnings"),
        "layers_direct": {k: dict(v) for k, v in sorted(direct.items())},
        "layers_in_blocks_expanded": {k: dict(v) for k, v in sorted(expanded.items())},
        "inserts_in_modelspace": n_inserts, "max_nested_block_depth": max_depth, "xrefs": xrefs,
        "texts_total": len(texts), "texts_by_type": dict(Counter(t["type"] for t in texts)),
        "label_texts": len(labels), "duplicate_labels_within_3m": dups,
        "rooms_after_dedup": len(ex["rooms"]),
        "rooms_with_boundary": sum(1 for r in ex["rooms"] if r["boundary"] == "drawn"),
        "rooms_without_boundary": sum(1 for r in ex["rooms"] if r["boundary"] != "drawn"),
        "rooms_by_category": dict(Counter(r["category"] for r in ex["rooms"])),
        "doors": len(ex.get("doors", [])), "furniture": len(ex.get("furniture", [])), "grids": len(ex.get("grids", [])),
    }
    info["label_list"] = [{k: t[k] for k in ("text", "type", "layer", "x_m", "y_m", "room_name_like", "room_number_like")}
                          for t in labels]
    return info, rooms_json


# ======================================================================= section 4 — raster analysis
class Raster:
    """Room-index raster of one floor (metres). 0 = outside every room."""

    def __init__(self, fl, margin_m: float):
        um = fl.unit_m
        xs = [p[0] * um for r in fl.rooms for p in r.poly]
        ys = [p[1] * um for r in fl.rooms for p in r.poly]
        self.bx0, self.by0 = min(xs) - margin_m, min(ys) - margin_m
        self.bx1, self.by1 = max(xs) + margin_m, max(ys) + margin_m
        span = max(self.bx1 - self.bx0, self.by1 - self.by0)
        self.res = max(0.1, span / 2500)
        W, H = int(math.ceil((self.bx1 - self.bx0) / self.res)), int(math.ceil((self.by1 - self.by0) / self.res))
        img = Image.new("I", (W, H), 0)
        dr = ImageDraw.Draw(img)
        self.idx = {}
        order = sorted(range(len(fl.rooms)), key=lambda i: -fl.rooms[i].area_m2)  # smallest room wins (Floor.room_at)
        for i in order:
            r = fl.rooms[i]
            self.idx[r.id] = i + 1
            dr.polygon([self.px(x * um, y * um) for x, y in r.poly], fill=i + 1)
        self.L = np.array(img, dtype=np.int32)
        self.H, self.W = self.L.shape
        # building outline = union of rooms, closed by 0.5 m (wall gaps) and with holes filled
        B = self.L > 0
        k = max(1, int(round(0.5 / self.res)))
        Bc = self._erode(self._dilate(B, k), k)
        im = Image.fromarray((Bc * 255).astype(np.uint8)).copy()  # copy: fromarray images share a read-only buffer
        ImageDraw.floodfill(im, (0, 0), 128)
        a = np.array(im)
        self.outline = a != 128
        self.edge = self.outline & ~self._erode(self.outline, 1)
        ey, ex_ = np.nonzero(self.edge)
        self.edge_xy = np.stack([self.bx0 + (ex_ + 0.5) * self.res, self.by1 - (ey + 0.5) * self.res], axis=1)

    @staticmethod
    def _dilate(M, k):
        out = M.copy()
        for _ in range(k):
            o = out.copy()
            o[1:, :] |= out[:-1, :]
            o[:-1, :] |= out[1:, :]
            o[:, 1:] |= out[:, :-1]
            o[:, :-1] |= out[:, 1:]
            out = o
        return out

    @staticmethod
    def _erode(M, k):
        return ~Raster._dilate(~M, k)

    def px(self, xm, ym):
        return ((xm - self.bx0) / self.res, (self.by1 - ym) / self.res)

    def ij(self, xm, ym):
        return int((self.by1 - ym) / self.res), int((xm - self.bx0) / self.res)

    def wedge(self, xm, ym, aim_deg, r_m, half_deg):
        """Boolean mask of the wedge in a window; returns (mask, (i0, i1, j0, j1))."""
        i0, j0 = self.ij(xm - r_m, ym + r_m)
        i1, j1 = self.ij(xm + r_m, ym - r_m)
        i0, j0 = max(0, i0), max(0, j0)
        i1, j1 = min(self.H, i1 + 1), min(self.W, j1 + 1)
        jj, ii = np.meshgrid(np.arange(j0, j1), np.arange(i0, i1))
        X = self.bx0 + (jj + 0.5) * self.res - xm
        Y = self.by1 - (ii + 0.5) * self.res - ym
        D = np.hypot(X, Y)
        A = np.degrees(np.arctan2(Y, X))
        diff = (A - aim_deg + 180) % 360 - 180
        return (D <= r_m) & (np.abs(diff) <= half_deg), (i0, i1, j0, j1)


def bearing_diff(a, b):
    return abs((a - b + 180) % 360 - 180)


def device_diagnostics(res, project, sel):
    pp = merged("cctv", project["systems"]["params"].get("cctv"))
    ppm = {"Identify": float(pp["ppm_identify"]), "Recognise": float(pp["ppm_recognise"]), "Observe": float(pp["ppm_observe"])}
    rooms_by = {r.id: r for f in res.floors for r in f.rooms}
    fl_by = {f.key: f for f in res.floors}
    devices, ranges = [], {}
    cam_stats, coverage, entries, overlaps = [], [], [], []
    for fl in res.floors:
        um = fl.unit_m
        fdev = [d for d in res.devices if d["floor"] == fl.key]
        cams = [d for d in fdev if d["sys"] == "cctv"]
        maxr = max([camera_range(d["item"], ppm.get(d.get("target") or "Recognise", 125)) for d in cams] or [10])
        R = Raster(fl, margin_m=maxr + 2)
        U_drawn = np.zeros_like(R.outline)
        U_wall = np.zeros_like(R.outline)
        for d in cams:
            tgt = d.get("target") or "Recognise"
            rad = round(camera_range(d["item"], ppm[tgt]), 1)  # same rounding as Backend.get_layout
            ranges[d["id"]] = rad
            xm, ym = d["x"] * um, d["y"] * um
            own = R.idx.get(d.get("room"), -1)
            out = {}
            for mode, half in (("app_wedge", WEDGE_HALF_DEG),
                               ("catalogue_hfov", min(catalog.CAMERA_OPTICS.get(d["item"], {"hfov": 90})["hfov"], 359) / 2)):
                M, (i0, i1, j0, j1) = R.wedge(xm, ym, d.get("a", 0) or 0, rad, half)
                Lw, Ow = R.L[i0:i1, j0:j1], R.outline[i0:i1, j0:j1]
                n = int(M.sum()) or 1
                out[mode] = {"area_m2": round(n * R.res ** 2, 1),
                             "pct_outside_building": round(100 * int((M & ~Ow).sum()) / n, 1),
                             "pct_crossing_wall": round(100 * int((M & Ow & (Lw != own)).sum()) / n, 1),
                             "pct_in_own_room": round(100 * int((M & (Lw == own)).sum()) / n, 1)}
                if mode == "app_wedge":
                    U_drawn[i0:i1, j0:j1] |= M
                    U_wall[i0:i1, j0:j1] |= M & (Lw == own)
            ext = None
            if len(R.edge_xy):
                dd = np.hypot(R.edge_xy[:, 0] - xm, R.edge_xy[:, 1] - ym)
                k = int(np.argmin(dd))
                brg = math.degrees(math.atan2(R.edge_xy[k, 1] - ym, R.edge_xy[k, 0] - xm))
                ext = {"dist_m": round(float(dd[k]), 2), "bearing_deg": round(brg, 1),
                       "aim_minus_bearing_deg": round(bearing_diff(d.get("a", 0) or 0, brg), 1)}
            cam_stats.append({"id": d["id"], "floor": fl.name, "item": d["item"], "target": tgt, "range_m": rad,
                              "room": rooms_by[d["room"]].name if d.get("room") in rooms_by else None,
                              "room_category": rooms_by[d["room"]].category if d.get("room") in rooms_by else None,
                              **{f"{m}_{k}": v for m, s in out.items() for k, v in s.items()},
                              "nearest_exterior_wall": ext,
                              "aimed_at_nearest_exterior_wall_lt45": bool(ext and ext["aim_minus_bearing_deg"] < 45)})
        # coverage of key spaces
        key_cats = {"lift_lobby", "lift", "stair", "entrance", "reception", "parking", "loading", "corridor"}
        for r in fl.rooms:
            if r.category in key_cats or re.search(r"RAMP|ENTRY|EXIT|AISLE|DRIVE", r.name, re.I):
                Rm = R.L == R.idx[r.id]
                n = int(Rm.sum()) or 1
                cd = round(100 * int((Rm & U_drawn).sum()) / n, 1)
                cw = round(100 * int((Rm & U_wall).sum()) / n, 1)
                coverage.append({"floor": fl.name, "room_id": r.id, "name": r.name, "category": r.category,
                                 "area_m2": round(r.area_m2, 1), "cameras_in_room": sum(1 for c in cams if c.get("room") == r.id),
                                 "pct_covered_as_drawn": cd, "pct_covered_wall_aware": cw, "covered(>=50% wall-aware)": cw >= 50})
        # entry points: doors serving entrance/lobby rooms or on the building outline
        for door in fl.doors:
            cats = {rooms_by[x].category for x in door.rooms if x in rooms_by}
            i, j = R.ij(door.pt[0] * um, door.pt[1] * um)
            on_edge = False
            if 0 <= i < R.H and 0 <= j < R.W:
                k = max(1, int(round(1.0 / R.res)))
                on_edge = bool(R.edge[max(0, i - k):i + k + 1, max(0, j - k):j + k + 1].any())
            if not (cats & {"entrance", "lift_lobby", "stair", "parking", "loading"} or on_edge):
                continue
            inside = 0 <= i < R.H and 0 <= j < R.W
            entries.append({"floor": fl.name, "door": door.id, "tag": door.tag, "inferred": door.inferred,
                            "rooms": [rooms_by[x].name for x in door.rooms if x in rooms_by], "on_building_outline": on_edge,
                            "covered_as_drawn": bool(inside and U_drawn[i, j]), "covered_wall_aware": bool(inside and U_wall[i, j])})
        # overlapping symbols within 0.5 m
        pts = [(d, d["x"] * um, d["y"] * um) for d in fdev]
        for a in range(len(pts)):
            for b in range(a + 1, len(pts)):
                dm = math.hypot(pts[a][1] - pts[b][1], pts[a][2] - pts[b][2])
                if dm < 0.5:
                    overlaps.append({"floor": fl.name, "a": pts[a][0]["id"], "a_item": pts[a][0]["item"],
                                     "b": pts[b][0]["id"], "b_item": pts[b][0]["item"], "dist_m": round(dm, 3),
                                     "room": rooms_by[pts[a][0]["room"]].name if pts[a][0].get("room") in rooms_by else None})
    for d in res.devices:
        fl = fl_by[d["floor"]]
        um = fl.unit_m
        r = rooms_by.get(d.get("room") or "")
        rt = res.routes.get(d["id"])
        row = {"id": d["id"], "system": d["sys"], "item": d["item"], "type": type_label(d), "floor": fl.name,
               "floor_key": fl.key, "x_m": round(d["x"] * um, 3), "y_m": round(d["y"] * um, 3),
               "room_id": d.get("room"), "room_name": r.name if r else None, "room_category": r.category if r else None,
               "room_boundary": r.boundary if r else None, "mount_type": None, "auto": d.get("auto", True),
               "edited": bool(d.get("edited")), "in_selected_systems": d["sys"] in sel}
        if d["sys"] == "cctv":
            o = catalog.CAMERA_OPTICS.get(d["item"], {"hres": 2560, "hfov": 90})
            tgt = d.get("target") or "Recognise"
            row.update({"aim_deg": d.get("a", 0), "hfov_deg": o["hfov"], "target": tgt, "fov_radius_m": ranges.get(d["id"]),
                        "fov_formula": "hres/(ppm*2*tan(min(hfov,120)/2))",
                        "fov_inputs": {"hres": o["hres"], "hfov_used": min(o["hfov"], 120), "ppm": ppm[tgt]},
                        "range_by_target_m": d.get("range_m"), "wedge_half_angle_drawn_deg": round(WEDGE_HALF_DEG, 2)})
        elif "a" in d:
            row["aim_deg"] = d["a"]
        for k in ("sub", "kind", "ports", "door", "tag", "leaves", "readers", "lock", "inferred", "size", "src", "note"):
            if k in d:
                row[k] = d[k]
        if rt:
            cr = rooms_by.get(rt["room"])
            row.update({"comms_room": cr.name if cr else rt["room"], "comms_role": cr.comms if cr else None,
                        "comms_floor": fl_by[rt["room_floor"]].name if rt["room_floor"] in fl_by else rt["room_floor"],
                        "cable_len_m": rt["len_m"], "over_90m": rt["over"], "riser": rt["riser"], "route_pts_m":
                            [[round(x * um, 2), round(y * um, 2)] for x, y in rt["pts"]]})
        else:
            row.update({"comms_room": None, "cable_len_m": None, "over_90m": None})
        devices.append(row)
    return devices, ranges, cam_stats, coverage, entries, overlaps


def device_tables(res, devices, cam_stats, sel):
    out = {}
    out["counts_by_system"] = dict(Counter(d["system"] for d in devices))
    out["counts_by_floor_system"] = {f: dict(Counter(d["system"] for d in devices if d["floor"] == f))
                                     for f in sorted({d["floor"] for d in devices})}
    out["counts_by_system_category"] = {s: dict(Counter(d["room_category"] for d in devices if d["system"] == s))
                                        for s in SYSTEM_KEYS}
    out["counts_by_item"] = dict(Counter(d["item"] for d in devices))
    per100 = {}
    for fl in res.floors:
        area = sum(r.area_m2 for r in fl.rooms)
        n = Counter(d["system"] for d in devices if d["floor_key"] == fl.key)
        per100[fl.name] = {"room_area_m2": round(area, 1), **{s: round(n.get(s, 0) / area * 100, 2) for s in SYSTEM_KEYS},
                           "all": round(sum(n.values()) / area * 100, 2)}
    out["devices_per_100m2"] = per100
    fov = {}
    for it in sorted({c["item"] for c in cam_stats}):
        cs = [c for c in cam_stats if c["item"] == it]
        rr = [c["range_m"] for c in cs]
        fov[it] = {"n": len(cs), "targets": dict(Counter(c["target"] for c in cs)), "radius_min_m": min(rr),
                   "radius_mean_m": round(statistics.mean(rr), 2), "radius_max_m": max(rr),
                   "hfov_catalogue_deg": catalog.CAMERA_OPTICS[it]["hfov"], "wedge_drawn_deg": round(2 * WEDGE_HALF_DEG, 1)}
    out["fov_radius_by_camera_type"] = fov

    def agg(key):
        v = [c[key] for c in cam_stats]
        return {"mean": round(statistics.mean(v), 1) if v else None, "median": round(statistics.median(v), 1) if v else None,
                "max": max(v) if v else None, "cameras_gt_10pct": sum(1 for x in v if x > 10),
                "cameras_gt_50pct": sum(1 for x in v if x > 50)}
    out["fov_outside_building_pct"] = {"app_wedge": agg("app_wedge_pct_outside_building"),
                                       "catalogue_hfov": agg("catalogue_hfov_pct_outside_building")}
    out["fov_crossing_wall_pct"] = {"app_wedge": agg("app_wedge_pct_crossing_wall"),
                                    "catalogue_hfov": agg("catalogue_hfov_pct_crossing_wall")}
    out["cameras_in_parking_bays"] = {"value": None, "reason": "parking bays are not extracted or modelled by the app",
                                      "cameras_in_parking_rooms": sum(1 for c in cam_stats if c["room_category"] == "parking")}
    out["cameras_aimed_at_nearest_exterior_wall_lt45"] = [c["id"] for c in cam_stats if c["aimed_at_nearest_exterior_wall_lt45"]]
    out["aisles_detected"] = {"value": 0, "reason": "the app has no drive-aisle detection; corridors are the only linear spaces"}
    lens = [d["cable_len_m"] for d in devices if d.get("cable_len_m") is not None and d["system"] in sel and d["system"] in LIMITED]
    bins = [(0, 30), (30, 60), (60, 75), (75, 90), (90, 120), (120, 10 ** 6)]
    out["cable_length_distribution_m"] = {
        "n": len(lens), "min": min(lens) if lens else None, "median": statistics.median(lens) if lens else None,
        "mean": round(statistics.mean(lens), 1) if lens else None, "max": max(lens) if lens else None,
        "bins": {f"{a}-{b if b < 10 ** 6 else 'inf'}": sum(1 for x in lens if a <= x < b) for a, b in bins}}
    out["links_over_90m"] = [{"id": d["id"], "system": d["system"], "floor": d["floor"], "room": d["room_name"],
                              "comms_room": d["comms_room"], "len_m": d["cable_len_m"], "riser": d.get("riser")}
                             for d in devices if d.get("over_90m")]
    return out


# ======================================================================= section 5 — comms
def comms_diag(res):
    rows = []
    for rs in res.sizing["rooms"]:
        poe_pct = round(rs.poe_load_w / rs.poe_budget_w * 100, 1) if rs.poe_budget_w else 0.0
        racks = [{"name": rk.name, "used_u": rk.used, "pct": round(rk.used / 42 * 100, 1),
                  "items": [{"label": i.label, "item": i.item, "size_u": i.size, "u": i.u} for i in rk.items]} for rk in rs.racks]
        u_total = len(rs.racks) * 42
        rows.append({
            "id": rs.id, "name": rs.name, "floor": rs.floor_name, "role": rs.role, "devices_served": rs.served,
            "devices_served_total": sum(rs.served.values()), "elv_ports": rs.elv_ports, "data_ports": rs.data_ports,
            "ap_count": rs.ap_count, "elv_switches": rs.elv_switches, "user_switches": rs.user_switches,
            "poe_load_w": round(rs.poe_load_w, 1), "poe_budget_w": rs.poe_budget_w, "poe_pct": poe_pct,
            "elv_panels": rs.elv_panels, "data_panels": rs.data_panels, "fibre_panels": rs.fibre_panels,
            "cable_mgrs": rs.cable_mgrs, "pdus": rs.pdus, "equipment": rs.equipment, "load_w": round(rs.load_w, 1),
            "ups": catalog.get(rs.ups).part_no if rs.ups else None, "ups_item": rs.ups, "ups_qty": rs.ups_qty,
            "battery_packs": rs.packs, "battery_packs_qty": rs.packs_qty, "runtime_min": round(rs.runtime_min, 1),
            "racks": racks, "u_used": rs.u_used, "u_total": u_total, "u_pct": round(rs.u_used / u_total * 100, 1) if u_total else 0,
            "flag_poe_gt_80": poe_pct > 80, "flag_runtime_lt_60": bool(rs.ups) and rs.runtime_min < 60,
            "flag_u_gt_70": any(rk["pct"] > 70 for rk in racks),
        })
    by_sys = {s: {"rooms": len(sz["rooms"]), "u_used": sum(r.u_used for r in sz["rooms"]),
                  "racks": sum(len(r.racks) for r in sz["rooms"]),
                  "ups": [f"{r.name}: {r.ups} x{r.ups_qty} +{r.packs_qty} packs, {r.runtime_min:.0f} min" for r in sz["rooms"] if r.ups]}
              for s, sz in res.sizing_by_sys.items()}
    return {"combined_scenario": rows, "standalone_scenarios": by_sys,
            "backbone": res.backbone}


# ======================================================================= section 6 — BOQ
def boq_diag(res, label):
    b = res.boq
    t = b["totals"]
    lines = all_lines(b, include_excluded=True)
    per_tab = {}
    for k, tab in b["tabs"].items():
        tl = [l for s in tab.sections for l in s.lines] + (tab.shared_section.lines if tab.shared_section else [])
        per_tab[k] = {"lines": len(tl), "lines_qty_gt0": sum(1 for l in tl if l.qty > 0),
                      "combined_total": str(tab.combined), "standalone_total": str(tab.standalone),
                      "budgetary_count": sum(1 for l in tl if l.budget), "budgetary_value": str(sum((l.total for l in tl if l.budget), Decimal(0))),
                      "unpriced": sum(1 for l in tl if l.unpriced)}
    issues = defaultdict(list)
    keys = Counter(l.key for l in lines)
    tab_items = Counter((l.tab, l.item, l.section, l.description) for l in lines if l.item)
    for l in lines:
        ref = {"key": l.key, "tab": l.tab, "section": l.section, "item": l.item, "desc": l.description, "qty": str(l.qty),
               "list": None if l.list_price is None else str(l.list_price), "cur": l.currency, "part": l.part_no}
        if l.qty == 0:
            issues["qty_zero"].append(ref)
        if l.qty < 0:
            issues["qty_negative"].append(ref)
        if l.qty != l.qty.to_integral_value():
            issues["qty_non_integer"].append(ref)
        if l.qty > 0 and l.list_price is None:
            issues["price_missing"].append(ref)
        if l.qty > 0 and l.list_price is not None and l.list_price == 0:
            issues["price_zero"].append(ref)
        if l.currency != "QAR" and l.currency not in res.fx:
            issues["currency_without_fx"].append(ref)
        for field in ("description", "manufacturer", "part_no", "note", "unit"):
            v = str(getattr(l, field) or "")
            for word, rx in BAD_WORDS:
                if rx.search(v):
                    issues[f"text_{word.replace(' ', '_')}"].append({**ref, "field": field, "value": v})
        if keys[l.key] > 1:
            issues["duplicate_key"].append(ref)
    dup_items = [{"tab": k[0], "item": k[1], "section": k[2], "n": n} for k, n in tab_items.items() if n > 1]
    return {
        "label": label, "combined_award": b["combined"], "selected": res.selected,
        "totals": {k: (str(v) if not isinstance(v, dict) else {kk: {k3: str(v3) for k3, v3 in vv.items()} for kk, vv in v.items()})
                   for k, v in t.items()},
        "totals_formatted": {k: fmt(v) for k, v in t.items() if not isinstance(v, dict)},
        "per_tab": per_tab, "unpriced_lines_in_scope": b["unpriced"], "rfq_budgetary_in_scope": b["rfq"],
        "lines_in_scope": b["lines_in_scope"], "conf_mix": b["conf_mix"],
        "issue_counts": {k: len(v) for k, v in issues.items()}, "issues": issues, "duplicate_item_in_section": dup_items,
        "fx_used": {k: str(v) for k, v in res.fx.items()},
    }


def checks_diag(res, label):
    ck = res.checks
    return {"label": label, "total": ck["total"], "passed": ck["passed"], "blocking": ck["blocking"],
            "warnings_failed": sum(1 for g in ck["groups"] for r in g["rows"] if not r["pass"] and r["sev"] == "warn"),
            "groups": ck["groups"]}


# ======================================================================= section 7 — outputs
def recalc_xlsx(files: list[Path], outdir: Path, log) -> dict:
    outdir.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")  # Excel COM is not used by this helper
    if not soffice:
        return {"method": None, "error": "LibreOffice not found", "files": {}}
    prof = outdir / "lo_profile"
    out = {"method": f"LibreOffice headless ({soffice}) convert-to xlsx (formulas recalculated on load)", "files": {}}
    for f in files:
        t0 = time.perf_counter()
        p = subprocess.run([soffice, f"-env:UserInstallation=file://{prof.resolve()}", "--headless", "--calc",
                            "--convert-to", "xlsx", "--outdir", str(outdir), str(f)],
                           capture_output=True, text=True, timeout=300)
        log.write(f"soffice {f.name}: rc={p.returncode} {p.stdout.strip()} {p.stderr.strip()[-400:]}\n")
        out["files"][f.name] = {"recalc_file": str(outdir / f.name), "rc": p.returncode, "seconds": round(time.perf_counter() - t0, 2)}
    return out


def scan_xlsx(orig: Path, recalc: Path) -> dict:
    import openpyxl
    wf = openpyxl.load_workbook(orig, data_only=False)
    wv = openpyxl.load_workbook(recalc, data_only=True)
    errors, n_formulas, no_value = [], 0, 0
    for ws in wf.worksheets:
        wsv = wv[ws.title] if ws.title in wv.sheetnames else None
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_formulas += 1
                    v = wsv[c.coordinate].value if wsv is not None else None
                    if v is None:
                        no_value += 1
                    if isinstance(v, str) and v.startswith(XL_ERRORS):
                        errors.append({"sheet": ws.title, "cell": c.coordinate, "formula": c.value, "value": v})
                elif isinstance(c.value, str) and c.value.startswith(XL_ERRORS):
                    errors.append({"sheet": ws.title, "cell": c.coordinate, "formula": None, "value": c.value})
    return {"sheets": wf.sheetnames, "formulas": n_formulas, "formulas_without_value_after_recalc": no_value,
            "errors": errors, "_wv": wv}


def quote_totals_check(scan: dict, res) -> list[dict]:
    wv = scan["_wv"]
    ws = wv["Summary"]
    t = res.boq["totals"]
    rows = {}
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
        a = r[0].value
        if isinstance(a, str) and a not in rows:
            rows[a] = r
    out = []

    def cmp(label, col, app):
        r = rows.get(label)
        got = r[col].value if r is not None else None
        ok = got is not None and abs(Decimal(str(got)) - app) < Decimal("0.005")
        out.append({"label": label, "excel": got, "app": str(app), "match": ok})
    for s in res.selected:
        cmp(system_name(s), 2, t["systems"][s]["combined"] if res.boq["combined"] else t["systems"][s]["standalone"])
    if res.boq["combined"]:
        cmp("Shared infrastructure (racks, UPS, backbone, containment)", 2, t["shared"])
    cmp("TOTAL BASE PRICE (QAR)", 2, t["base"])
    if "network" in res.boq["tabs"]:
        cmp("Network Active equipment (quoted separately, not in base total)", 2, t["network"])
    if res.boq["combined"]:
        cmp("Saving with combined award", 3, t["saving"])
    return out


def costing_totals_check(scan: dict, res) -> list[dict]:
    ws = scan["_wv"]["Summary"]
    out = []
    for r in ws.iter_rows(min_row=4, max_row=ws.max_row):
        if r[0].value is None:
            continue
        out.append({"tab": r[0].value, "sell": r[1].value, "cost": r[2].value, "app_total": r[5].value, "difference": r[6].value,
                    "match": r[6].value is not None and abs(float(r[6].value)) < 0.005})
    return out


def scan_pdf(path: Path, expect: dict[str, str]) -> dict:
    import pymupdf as fitz
    doc = fitz.open(str(path))
    hits = []
    found = {k: [] for k in expect}
    for i, page in enumerate(doc):
        txt = page.get_text()
        for word, rx in PDF_TERMS:
            for m in rx.finditer(txt):
                a, b = max(0, m.start() - 50), min(len(txt), m.end() + 50)
                hits.append({"term": word, "page": i + 1, "context": " ".join(txt[a:b].split())})
        for k, s in expect.items():
            if s in txt:
                found[k].append(i + 1)
    n = len(doc)
    doc.close()
    return {"pages": n, "term_hits": hits, "term_counts": dict(Counter(h["term"] for h in hits)),
            "expected_amounts": {k: {"text": v, "found_on_pages": found[k]} for k, v in expect.items()}}


def render_pdf_page(path: Path, out_png: Path, dpi=150):
    import pymupdf as fitz
    doc = fitz.open(str(path))
    pix = doc[0].get_pixmap(dpi=dpi, alpha=False)
    out_png.parent.mkdir(parents=True, exist_ok=True)
    pix.save(str(out_png))
    doc.close()
    return {"png": out_png.name, "width": pix.width, "height": pix.height, "dpi": dpi}


def outputs_diag(res, project, out_folder: Path, files: list[str], dest: Path, recalc_dir: Path, img_dir: Path, tag, log):
    t = res.boq["totals"]
    rep = {"folder": str(out_folder), "files": {}}
    dest.mkdir(parents=True, exist_ok=True)
    for f in files:
        shutil.copy2(out_folder / f, dest / f)
        rep["files"][f] = {"bytes": (out_folder / f).stat().st_size}
    xl = [out_folder / f for f in files if f.endswith(".xlsx")]
    with timed(f"{tag}: Excel recalc (LibreOffice)"):
        rc = recalc_xlsx(xl, recalc_dir, log)
    rep["excel_recalc"] = {k: v for k, v in rc.items()}
    for f in xl:
        rf = recalc_dir / f.name
        if not rf.exists():
            rep["files"][f.name]["scan"] = {"error": "recalculated file missing"}
            continue
        sc = scan_xlsx(f, rf)
        if "Quotation" in f.name:
            sc["totals_vs_app"] = quote_totals_check(sc, res)
        if "Internal_Costing" in f.name:
            sc["totals_vs_app"] = costing_totals_check(sc, res)
        sc.pop("_wv", None)
        rep["files"][f.name]["scan"] = sc
    for f in files:
        if not f.endswith(".pdf"):
            continue
        exp = {}
        if f.endswith("_Quotation.pdf") and "_Rev" in f and not any(x in f for x in ("_CCTV_", "_ACS_", "_IPTV_", "_SC_", "_IC_", "_INT_")):
            exp = {"base": fmt(t["base"]), "shared": fmt(t["shared"]), "network": fmt(t["network"]), "saving": fmt(t["saving"])}
            for s in res.selected:
                exp[f"{s}_combined"] = fmt(t["systems"][s]["combined"])
        elif "Combined_Summary" in f:
            exp = {"base": fmt(t["base"]), "standalone": fmt(t["standalone"]), "saving": fmt(t["saving"])}
        else:
            for s, short in (("cctv", "CCTV"), ("access", "ACS"), ("iptv", "IPTV"), ("cabling", "SC"), ("intercom", "IC"), ("intrusion", "INT")):
                if f"_{short}_Quotation" in f and s in t["systems"]:
                    exp = {f"{s}_standalone": fmt(t["systems"][s]["standalone"]), f"{s}_combined": fmt(t["systems"][s]["combined"])}
        rep["files"][f]["scan"] = scan_pdf(out_folder / f, exp)
        if f.endswith("_Layouts.pdf"):
            with timed(f"{tag}: render layout PDF page 1"):
                rep["files"][f]["page1_png"] = render_pdf_page(out_folder / f, img_dir / f"{tag}_layout_pdf_page1_150dpi.png")
    return rep


# ======================================================================= main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out_dir")
    ap.add_argument("--label", default="S")
    ap.add_argument("--drawings", nargs="*")
    ap.add_argument("--name", default="Sample HQ Fit-Out (diagnostics)")
    a = ap.parse_args()
    OUT = Path(a.out_dir).resolve()
    tag = a.label
    work = OUT / "work" / tag
    if work.exists():
        shutil.rmtree(work)
    for sub in ("json", "images", "logs", "outputs", "inputs"):
        (OUT / sub).mkdir(parents=True, exist_ok=True)
    log = open(OUT / "logs" / f"backend_calls_{tag}.log", "w", encoding="utf-8")
    t_all = time.perf_counter()

    # ---------------------------------------------------------------- inputs
    with timed(f"{tag}: prepare inputs"):
        if a.drawings:
            files = [str(Path(f).resolve()) for f in a.drawings]
        else:
            from elv_engine.samples import make_set
            files = make_set(str(OUT / "inputs" / tag))
    with timed(f"{tag}: Backend start"):
        b = Backend(data_dir=work / "data", projects_dir=work / "projects")
    call = Client(b, log)
    call("create_project", {"name": a.name, "client": "Diagnostics", "scope": tag})
    with timed(f"{tag}: import + extraction (backend job, all drawings)"):
        job = call("import_drawings", {"paths": files})["job"]
        call.wait(job)
    st = call("project_state")
    p = b.project
    failed = [d for d in p["drawings"] if d["status"] == "Failed"]
    # per-drawing extraction timing (same function the job uses, separate cache)
    from elv_engine.extract import extract_drawing
    for d in p["drawings"]:
        src = b.folder / "drawings" / d["file"]
        with timed(f"{tag}: extract_drawing {d['name']} ({d['kind']}, {src.stat().st_size // 1024} KB)"):
            try:
                extract_drawing(str(src), str(work / "timing_cache"), d["id"] + "t", d.get("page", 0))
            except Exception as exc:  # noqa: BLE001
                log.write(f"timing extract failed {d['name']}: {exc}\n")

    # ---------------------------------------------------------------- section 3
    draw_rep, rooms_all = [], []
    for d in p["drawings"]:
        if not d.get("extraction"):
            draw_rep.append({"file": d["file"], "status": d["status"], "error": d["error"]})
            continue
        ex = d["extraction"]
        unit_m = (d.get("calibration") or {}).get("unit_m") or ex.get("unit_m") or 0.001
        src = b.folder / "drawings" / d["file"]
        dxf = src if d["kind"] == "dxf" else b.folder / "cache" / f"{d['id']}.dxf"
        if d["kind"] == "pdf":
            draw_rep.append({"file": d["file"], "kind": "pdf", "note": "PDF: layer/entity statistics not applicable",
                             "unit_m": unit_m, "rooms": len(ex["rooms"])})
            continue
        with timed(f"{tag}: drawing diagnostics {d['name']}"):
            info, rooms = drawing_diag(src, dxf, ex, unit_m, d)
        info["status"] = d["status"]
        info["floor"] = d["floor"]
        draw_rep.append(info)
        rooms_all += rooms
    jdump(OUT / "json" / f"drawings_{tag}.json", draw_rep)
    jdump(OUT / "json" / f"rooms_{tag}.json", rooms_all)

    # ---------------------------------------------------------------- pipeline stage timings (same calls as pipeline.compute)
    from elv_engine import boq as boq_mod
    from elv_engine import checks as checks_mod
    from elv_engine import pipeline, quantities
    from elv_engine.layout import apply_edits
    from elv_engine.layout.floor import build_floors
    from elv_engine.layout.place import place_all
    from elv_engine.layout.routing import assign
    lib, fx = b.lib.as_map(), b._fx()
    params = p["systems"]["params"]
    sel = p["systems"]["selected"]
    with timed(f"{tag}: engine build_floors"):
        floors = build_floors(p)
    with timed(f"{tag}: engine place_all (placement)"):
        auto = place_all(floors, params)
    devs = apply_edits(auto, p.get("layout_edits"))
    with timed(f"{tag}: engine routing.assign"):
        rt = assign(floors, devs, merged("cabling", params.get("cabling")))
    with timed(f"{tag}: engine quantities.compute"):
        qres = quantities.compute(devs, rt["routes"], params)
    with timed(f"{tag}: engine comms sizing (all scenarios, infra_for)"):
        for s in sel:
            pipeline.infra_for({s}, floors, devs, rt["routes"], rt["backbone"], qres, params)
        pipeline.infra_for(set(sel), floors, devs, rt["routes"], rt["backbone"], qres, params)
    with timed(f"{tag}: engine pipeline.compute (complete, unpriced library)"):
        res0 = pipeline.compute(p, lib, fx)
    with timed(f"{tag}: engine checks.run"):
        checks_mod.run({"boq": res0.boq, "selected": res0.selected, "qinfo": res0.qres["info"], "devices": res0.devices,
                        "routes": res0.routes, "floors": res0.floors, "sizing": res0.sizing, "fx": fx, "commercial": res0.commercial})
    _ = boq_mod

    # ---------------------------------------------------------------- view models via the API (timed)
    with timed(f"{tag}: API get_drawings"):
        call("get_drawings")
    for fl in res0.floors:
        with timed(f"{tag}: API get_layout {fl.name}"):
            call("get_layout", {"floor": fl.key})
    with timed(f"{tag}: API get_comms"):
        call("get_comms")
    with timed(f"{tag}: API get_boq (unpriced)"):
        call("get_boq")
    with timed(f"{tag}: API get_checks (unpriced)"):
        call("get_checks")

    res = b.result()
    # ---------------------------------------------------------------- section 4
    with timed(f"{tag}: device diagnostics (raster FOV analysis)"):
        devices, ranges, cam_stats, coverage, entries, overlaps = device_diagnostics(res, p, set(sel))
        tables = device_tables(res, devices, cam_stats, set(sel))
    jdump(OUT / "json" / f"devices_{tag}.json", devices)
    jdump(OUT / "json" / f"cameras_fov_{tag}.json", cam_stats)
    jdump(OUT / "json" / f"device_tables_{tag}.json", {**tables, "key_space_coverage": coverage, "entry_points": entries,
                                                       "overlapping_symbols_lt_0_5m": overlaps})
    # ---------------------------------------------------------------- section 5
    jdump(OUT / "json" / f"comms_{tag}.json", comms_diag(res))
    # ---------------------------------------------------------------- section 6 (as shipped: empty price library)
    boq_unpriced = boq_diag(res, "as shipped (price library empty)")
    checks_unpriced = checks_diag(res, "as shipped (price library empty)")

    # ---------------------------------------------------------------- apply deterministic test prices to the WORK library
    from tests.conftest import TEST_PRICES
    with timed(f"{tag}: apply TEST_PRICES via update_price API ({len(b.lib.rows())} items)"):
        for row in call("get_prices")["rows"]:
            lp, cur, disc, conf = TEST_PRICES[row["category"]]
            call("update_price", {"key": row["key"], "patch": {"list_price": lp, "currency": cur, "discount": disc,
                                                                "confidence": conf}})
        call("update_price", {"key": "iptv_headend", "patch": {"confidence": "B"}})  # exercise the Budgetary path
    with timed(f"{tag}: engine pipeline.compute (priced)"):
        resP = b.result()
    boq_priced = boq_diag(resP, "TEST_PRICES (tests/conftest.py) + iptv_headend marked Budgetary")
    checks_priced = checks_diag(resP, "TEST_PRICES")
    jdump(OUT / "json" / f"boq_totals_{tag}.json", {"unpriced": boq_unpriced, "priced": boq_priced})
    jdump(OUT / "json" / f"checks_{tag}.json", {"unpriced": checks_unpriced, "priced": checks_priced})
    with timed(f"{tag}: API get_boq (priced)"):
        call("get_boq")

    # golden regression (repository golden = tests/golden/sample_totals.json, uses TEST_PRICES without the Budgetary flag)
    reg = None
    if not a.drawings:
        from tests.test_pipeline import summary
        call("update_price", {"key": "iptv_headend", "patch": {"confidence": "H"}})
        got = summary(b.result())
        want = json.loads((ROOT / "tests" / "golden" / "sample_totals.json").read_text(encoding="utf-8"))
        rows = []

        def walk(prefix, w, g):
            if isinstance(w, dict):
                for k in w:
                    walk(f"{prefix}.{k}" if prefix else k, w[k], (g or {}).get(k) if isinstance(g, dict) else None)
            else:
                diff = None
                try:
                    diff = str(Decimal(str(g)) - Decimal(str(w)))
                except Exception:  # noqa: BLE001
                    pass
                rows.append({"field": prefix, "expected": w, "actual": g, "difference": diff, "match": w == g})
        walk("", want, got)
        reg = {"golden_file": "tests/golden/sample_totals.json", "rows": rows, "all_match": all(r["match"] for r in rows)}
        jdump(OUT / "json" / f"regression_golden_{tag}.json", reg)
        call("update_price", {"key": "iptv_headend", "patch": {"confidence": "B"}})
    resP = b.result()

    # ---------------------------------------------------------------- section 7 — generate every output (as the Checks screen does)
    gen = {}
    call("set_export_options", {"selected": {k: True for k in ("full", "per", "summary", "layout", "racks", "cost", "zip")}})
    ck = call("get_checks")
    if ck["blocking"]:
        call("set_override", {"reason": f"TEST-ONLY diagnostics run: {ck['blocking']} blocking check(s) recorded in checks_{tag}.json"})
    for rev_tag, company in (("defaults", None), ("company", {"name": "Diagnostics ELV Contractor W.L.L.", "address": "P.O. Box 1, Doha, Qatar",
                                                              "phone": "+974 0000 0000", "email": "test@example.qa", "cr": "000000"})):
        if company:
            call("update_settings", {"company": company, "user": {"name": "Test Runner", "role": "ELV Estimator"}})
        with timed(f"{tag}: generate all outputs ({rev_tag}) — backend export job"):
            r = call("generate", {"bump": True, "change_log": f"diagnostics {rev_tag}"})
            j = call.wait(r["job"])
        res_now = b.result()
        o = j["result"]
        gen[rev_tag] = outputs_diag(res_now, p, Path(o["folder"]), o["files"], OUT / "outputs" / tag / f"Rev{p['revision']['current']}_{rev_tag}",
                                    work / f"recalc_{rev_tag}", OUT / "images", f"{tag}_{rev_tag}", log)
    # individual output builders (timing only, separate folder; no revision bump)
    from elv_engine.outputs import package
    import copy
    for k in ("full", "per", "summary", "layout", "racks", "cost"):
        with timed(f"{tag}: output builder '{k}' alone"):
            package.generate(copy.deepcopy(p), res_now, work / f"timing_{k}", {k}, b.settings["company"], b.settings["user"])
    jdump(OUT / "json" / f"outputs_{tag}.json", gen)

    # ---------------------------------------------------------------- section 9 — images
    from diag_images import render_floor
    img_files = []
    with timed(f"{tag}: diagnostic images (all floors)"):
        for fl in res_now.floors:
            lw = b.folder / "cache" / f"{fl.key}.lines.json.gz"
            fdev = [d for d in res_now.devices if d["floor"] == fl.key]
            img_files += render_floor(OUT / "images", tag, fl, fdev, res_now.routes, set(sel), lw, ranges)
    TIMES[f"{tag}: TOTAL diagnostics run"] = round(time.perf_counter() - t_all, 2)
    summary_out = {
        "label": tag, "inputs": files, "drawings_failed": [{"file": d["file"], "error": d["error"]} for d in failed],
        "header": st.get("header"), "floors": [{"key": f.key, "name": f.name, "rooms": len(f.rooms), "unit_m": f.unit_m,
                                                "calibrated": f.calibrated} for f in res_now.floors],
        "devices_total": len(res_now.devices), "checks_unpriced": f"{checks_unpriced['passed']}/{checks_unpriced['total']} "
                                                                  f"({checks_unpriced['blocking']} blocking)",
        "checks_priced": f"{checks_priced['passed']}/{checks_priced['total']} ({checks_priced['blocking']} blocking)",
        "regression_all_match": reg["all_match"] if reg else None, "images": [Path(x).name for x in img_files],
        "timings_s": TIMES,
    }
    jdump(OUT / "json" / f"run_summary_{tag}.json", summary_out)
    jdump(OUT / "json" / f"timings_{tag}.json", TIMES)
    log.close()
    print(json.dumps({k: v for k, v in summary_out.items() if k != "images"}, indent=1, default=str))


if __name__ == "__main__":
    main()
