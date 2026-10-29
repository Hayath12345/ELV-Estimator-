"""Internal costing workbook — full price build-up as live formulas.

Inputs sheet holds markup, labour markup, duty, freight and the exchange-rate table
(all popular currencies). Build-up sheet: one row per BOQ line with
FX lookup -> landed unit cost -> unit sell (ROUND to 2 dp) -> totals -> margin.
Summary uses SUMIFS per tab and reconciles against the app's totals.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from ..boq import all_lines
from ..currencies import CODES, NAMES
from ..money import dec
from ..systems import TAB_ORDER, tab_name
from .common import rev_label

TH = PatternFill("solid", fgColor="D9E1F2")
INP = PatternFill("solid", fgColor="FFF4D6")
MONEY = '#,##0.00;[Red]-#,##0.00'
CONF = {"H": "High", "M": "Medium", "B": "Budgetary", "O": "Other", "L": "Labour rate"}


def build(path: str, project: dict, res, checks_override: dict | None = None) -> None:
    boq = res.boq
    comm = res.commercial
    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    inp["A1"] = f"Internal costing — {project['info'].get('name', '')} — {project['info'].get('ref', '')} {rev_label(project)}"
    inp["A1"].font = Font(bold=True, size=12, color="14254A")
    inp["A2"] = "INTERNAL — NOT FOR ISSUE TO CLIENT. Yellow cells are inputs; every total recalculates."
    inp["A2"].font = Font(italic=True, size=9, color="B42318")
    rows = [("Markup", float(dec(comm.get("markup"), 0)) / 100),
            ("LabourMarkup", float(dec(comm.get("labour_markup") or comm.get("markup"), 0)) / 100),
            ("Duty", float(dec(comm.get("duty"), 0)) / 100),
            ("Freight", float(dec(comm.get("freight"), 0)) / 100)]
    labels = {"Markup": "Markup on materials", "LabourMarkup": "Markup on labour", "Duty": "Customs duty (imported items)",
              "Freight": "Freight & clearing (imported items)"}
    for i, (name, val) in enumerate(rows, 4):
        inp.cell(row=i, column=1, value=labels[name])
        c = inp.cell(row=i, column=2, value=val)
        c.number_format = "0.00%"
        c.fill = INP
        wb.defined_names[name] = DefinedName(name, attr_text=f"Inputs!$B${i}")
    inp.cell(row=10, column=1, value="Exchange rates").font = Font(bold=True, color="14254A")
    for c, h in enumerate(["Code", "Currency", "QAR per unit"], 1):
        cell = inp.cell(row=11, column=c, value=h)
        cell.font = Font(bold=True, size=9)
        cell.fill = TH
    r = 12
    for code in CODES:
        inp.cell(row=r, column=1, value=code)
        inp.cell(row=r, column=2, value=NAMES.get(code, code))
        c = inp.cell(row=r, column=3, value=float(res.fx.get(code, 1) or 1))
        c.number_format = "0.000000"
        c.fill = INP
        r += 1
    last_fx = r - 1
    wb.defined_names["FXTable"] = DefinedName("FXTable", attr_text=f"Inputs!$A$12:$C${last_fx}")
    wb.defined_names["FXCodes"] = DefinedName("FXCodes", attr_text=f"Inputs!$A$12:$A${last_fx}")
    inp.column_dimensions["A"].width = 36
    inp.column_dimensions["B"].width = 26
    inp.column_dimensions["C"].width = 16
    if checks_override:
        inp.cell(row=last_fx + 2, column=1, value="Issued with failing checks — override reason:").font = Font(bold=True, color="B42318")
        inp.cell(row=last_fx + 3, column=1, value=f"{checks_override.get('reason')} ({checks_override.get('by')}, {checks_override.get('at')})")

    ws = wb.create_sheet("Build-up")
    heads = ["Tab", "Section", "Description", "Manufacturer", "Part No.", "Category", "Unit", "Qty", "List price", "Curr.",
             "Disc. %", "FX to QAR", "Unit cost QAR", "Unit sell QAR", "Total cost", "Total sell", "Margin QAR", "Margin %",
             "Confidence", "Source", "Date checked", "In base"]
    widths = [18, 26, 48, 18, 18, 12, 6, 9, 11, 7, 7, 10, 13, 13, 14, 14, 13, 9, 11, 26, 12, 8]
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, size=9, color="14254A")
        cell.fill = TH
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "D2"
    r = 2
    lines = all_lines(boq, include_excluded=True)
    for l in lines:
        in_base = l.tab not in ("network", "optional") and not (l.shared and boq["combined"])
        vals = [tab_name(l.tab), ("Shared-eligible items" if l.shared else l.section), l.description, l.manufacturer,
                l.part_no, l.category, l.unit, float(l.qty), float(l.list_price) if l.list_price is not None else None,
                l.currency, float(l.discount)]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=12, value=f'=IF(J{r}="QAR",1,VLOOKUP(J{r},FXTable,3,FALSE))')
        ws.cell(row=r, column=13, value=f'=IF(I{r}="","",IF(F{r}="Labour",I{r}*L{r},IF(J{r}="QAR",I{r}*(1-K{r}/100),'
                                       f'I{r}*L{r}*(1-K{r}/100)*(1+Duty)*(1+Freight))))')
        ws.cell(row=r, column=14, value=f'=IF(M{r}="","",IF(F{r}="Labour",ROUND(M{r}*(1+LabourMarkup),2),ROUND(M{r}*(1+Markup),2)))')
        ws.cell(row=r, column=15, value=f'=IF(M{r}="",0,ROUND(H{r}*M{r},2))')
        ws.cell(row=r, column=16, value=f'=IF(N{r}="",0,ROUND(H{r}*N{r},2))')
        ws.cell(row=r, column=17, value=f"=P{r}-O{r}")
        ws.cell(row=r, column=18, value=f"=IF(P{r}=0,0,Q{r}/P{r})")
        ws.cell(row=r, column=19, value=CONF.get(l.confidence or "", "No price" if l.list_price is None else "Other"))
        ws.cell(row=r, column=20, value=l.source)
        ws.cell(row=r, column=21, value=l.date_checked)
        ws.cell(row=r, column=22, value="Y" if in_base else "N")
        for c in (9, 13, 14, 15, 16, 17):
            ws.cell(row=r, column=c).number_format = MONEY
        ws.cell(row=r, column=12).number_format = "0.0000"
        ws.cell(row=r, column=18).number_format = "0.0%"
        r += 1
    last = r - 1
    dv = DataValidation(type="list", formula1="=FXCodes", allow_blank=False)
    ws.add_data_validation(dv)
    if last >= 2:
        dv.add(f"J2:J{last}")
        ws.conditional_formatting.add(f"A2:V{last}", FormulaRule(formula=[f'$I2=""'], fill=PatternFill("solid", fgColor="FDECEC")))
        ws.conditional_formatting.add(f"A2:V{last}", FormulaRule(formula=[f'$S2="Budgetary"'], fill=PatternFill("solid", fgColor="FFF7E0")))
    ws.auto_filter.ref = f"A1:V{max(last, 1)}"

    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = f"Cost & margin summary — {project['info'].get('ref', '')} {rev_label(project)}"
    sm["A1"].font = Font(bold=True, size=12, color="14254A")
    for c, h in enumerate(["Tab", "Sell (QAR)", "Cost (QAR)", "Margin (QAR)", "Margin %", "App total (QAR)", "Difference"], 1):
        cell = sm.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True, size=9)
        cell.fill = TH
    rr = 4
    first = rr
    t = boq["totals"]
    for key in TAB_ORDER:
        if key not in boq["tabs"]:
            continue
        name = tab_name(key)
        sm.cell(row=rr, column=1, value=name + (" (separate)" if key == "network" else " (not in base)" if key == "optional" else ""))
        crit = '"Y"' if key not in ("network", "optional") else '"N"'
        sm.cell(row=rr, column=2, value=f"=SUMIFS('Build-up'!P:P,'Build-up'!A:A,\"{name}\",'Build-up'!V:V,{crit})")
        sm.cell(row=rr, column=3, value=f"=SUMIFS('Build-up'!O:O,'Build-up'!A:A,\"{name}\",'Build-up'!V:V,{crit})")
        sm.cell(row=rr, column=4, value=f"=B{rr}-C{rr}")
        sm.cell(row=rr, column=5, value=f"=IF(B{rr}=0,0,D{rr}/B{rr})")
        tab = boq["tabs"][key]
        app = tab.combined if (boq["combined"] or key not in res.selected) else tab.standalone
        sm.cell(row=rr, column=6, value=float(app))
        sm.cell(row=rr, column=7, value=f"=ROUND(B{rr}-F{rr},2)")
        for c in (2, 3, 4, 6, 7):
            sm.cell(row=rr, column=c).number_format = MONEY
        sm.cell(row=rr, column=5).number_format = "0.0%"
        rr += 1
    sm.cell(row=rr, column=1, value="Base total (excl. network & optional)").font = Font(bold=True)
    base_rows = [i for i in range(first, rr) if not str(sm.cell(row=i, column=1).value).endswith(("(separate)", "(not in base)"))]
    sm.cell(row=rr, column=2, value="=" + "+".join(f"B{i}" for i in base_rows) if base_rows else 0).number_format = MONEY
    sm.cell(row=rr, column=3, value="=" + "+".join(f"C{i}" for i in base_rows) if base_rows else 0).number_format = MONEY
    sm.cell(row=rr, column=4, value=f"=B{rr}-C{rr}").number_format = MONEY
    sm.cell(row=rr, column=5, value=f"=IF(B{rr}=0,0,D{rr}/B{rr})").number_format = "0.0%"
    sm.cell(row=rr, column=6, value=float(t["base"])).number_format = MONEY
    sm.cell(row=rr, column=7, value=f"=ROUND(B{rr}-F{rr},2)").number_format = MONEY
    sm.column_dimensions["A"].width = 40
    for col in "BCDEFG":
        sm.column_dimensions[col].width = 16

    ck = wb.create_sheet("Checks")
    ck["A1"] = f"Reconciliation checks — {res.checks['passed']} / {res.checks['total']} passed"
    ck["A1"].font = Font(bold=True, color="14254A")
    rr = 3
    for c, h in enumerate(["Group", "Check", "Value", "Rule", "Limit", "Result"], 1):
        ck.cell(row=rr, column=c, value=h).fill = TH
    rr += 1
    for g in res.checks["groups"]:
        for row in g["rows"]:
            vals = [g["g"], row["name"], row["a"], row["op"], row["b"], "PASS" if row["pass"] else ("WARN" if row["sev"] == "warn" else "FAIL")]
            for c, v in enumerate(vals, 1):
                ck.cell(row=rr, column=c, value=v)
            rr += 1
    ck.column_dimensions["A"].width = 22
    ck.column_dimensions["B"].width = 50
    ck.column_dimensions["C"].width = 18
    ck.column_dimensions["E"].width = 18
    wb.save(path)
